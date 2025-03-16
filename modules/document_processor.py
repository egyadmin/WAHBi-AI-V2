import os
import re
import io
import tempfile
from typing import Dict, List, Any, Union, Tuple, Optional
import pandas as pd
import numpy as np
from datetime import datetime

# المكتبات الخاصة بمعالجة أنواع المستندات المختلفة
import docx
import PyPDF2
import fitz  # PyMuPDF
import pdfplumber
import mammoth
from openpyxl import load_workbook
from PIL import Image
import pytesseract

# المكتبات الخاصة بالمعالجة الطبيعية للغة
import nltk
from nltk.tokenize import sent_tokenize, word_tokenize
from nltk.corpus import stopwords
from nltk import ngrams

# تحميل الموارد اللازمة للغة العربية
try:
    nltk.data.find('tokenizers/punkt')
    nltk.data.find('corpora/stopwords')
except LookupError:
    nltk.download('punkt')
    nltk.download('stopwords')

class DocumentProcessor:
    """
    فئة لمعالجة المستندات المختلفة وتحليلها واستخراج المعلومات منها
    تدعم الملفات بصيغة PDF, DOCX, XLSX, CSV, TXT
    """
    
    def __init__(self):
        """
        تهيئة معالج المستندات
        """
        # تحميل قائمة الكلمات الدلالية للمناقصات
        self.tender_keywords = self._load_tender_keywords()
        
        # تحميل قائمة المتطلبات الشائعة
        self.common_requirements = self._load_common_requirements()
        
        # الكلمات التوقفية في اللغة العربية
        self.arabic_stopwords = set(stopwords.words('arabic'))
        
    def process_document(self, file_content: bytes, file_extension: str, file_name: str) -> Dict[str, Any]:
        """
        معالجة المستند وتحليله حسب نوعه
        """
        with tempfile.NamedTemporaryFile(suffix=f".{file_extension}", delete=False) as temp_file:
            temp_file.write(file_content)
            temp_path = temp_file.name
        
        try:
            if file_extension.lower() == 'pdf':
                extracted_data = self._process_pdf(temp_path)
            elif file_extension.lower() in ['docx', 'doc']:
                extracted_data = self._process_docx(temp_path)
            elif file_extension.lower() in ['xlsx', 'xls']:
                extracted_data = self._process_excel(temp_path)
            elif file_extension.lower() == 'csv':
                extracted_data = self._process_csv(temp_path)
            elif file_extension.lower() == 'txt':
                extracted_data = self._process_txt(temp_path)
            else:
                extracted_data = {"error": f"نوع الملف {file_extension} غير مدعوم"}
            
            extracted_data["file_name"] = file_name
            extracted_data["file_type"] = file_extension
            extracted_data["processed_time"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            
            return extracted_data
            
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)
    
    def _process_pdf(self, file_path: str) -> Dict[str, Any]:
        """
        معالجة ملف PDF واستخراج النص والبيانات منه
        """
        extracted_data = {"text": "", "metadata": {}, "images": [], "tables": [], "pages": []}
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    extracted_text = page.extract_text()
                    if extracted_text:
                        extracted_data["text"] += extracted_text + "\n"
            
            if not extracted_data["text"].strip():
                extracted_data["text"] = self._apply_ocr_to_pdf(file_path)
        except Exception as e:
            extracted_data["error"] = f"خطأ في معالجة ملف PDF: {str(e)}"
        return extracted_data
    
    def _apply_ocr_to_pdf(self, file_path: str) -> str:
        """
        تطبيق OCR على ملف PDF لاستخراج النص من الصور
        """
        try:
            doc = fitz.open(file_path)
            ocr_text = ""
            for page in doc:
                pix = page.get_pixmap()
                img_data = pix.tobytes("png")
                with io.BytesIO(img_data) as img_stream:
                    img = Image.open(img_stream)
                    ocr_text += pytesseract.image_to_string(img, lang='ara+eng') + "\n"
            return ocr_text
        except Exception as e:
            return f"خطأ في OCR: {str(e)}"
    
    def _process_docx(self, file_path: str) -> Dict[str, Any]:
        """
        معالجة ملف Word (DOCX) واستخراج النص والبيانات منه
        """
        extracted_data = {"text": "", "metadata": {}, "images": [], "tables": [], "paragraphs": []}
        try:
            doc = docx.Document(file_path)
            extracted_data["text"] = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            if not extracted_data["text"].strip():
                with open(file_path, "rb") as docx_file:
                    result = mammoth.extract_raw_text(docx_file)
                    extracted_data["text"] = result.value
        except Exception as e:
            extracted_data["error"] = f"خطأ في معالجة ملف DOCX: {str(e)}"
        return extracted_data